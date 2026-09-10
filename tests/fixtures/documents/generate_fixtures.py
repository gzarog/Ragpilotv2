"""Regenerates the binary document fixtures in this directory.

Run with ``python tests/fixtures/documents/generate_fixtures.py`` from a
venv with the ``dev`` extras installed (python-docx/python-pptx/openpyxl --
all dev/test-only, see pyproject.toml). Not imported by the test suite
itself; the fixtures it writes are committed so tests don't depend on
these libraries' exact output being stable across versions.
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent


def _make_docx() -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Doc Title", level=1)
    doc.add_paragraph("Intro paragraph text.")
    doc.add_heading("Section One", level=2)
    doc.add_paragraph("Paragraph in section one.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "1"
    doc.save(HERE / "document.docx")


def _make_pptx() -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Presentation Title"
    body = slide.placeholders[1].text_frame
    body.text = "First bullet point"
    body.add_paragraph().text = "Second bullet point"
    prs.save(HERE / "presentation.pptx")


def _make_xlsx() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Name", "Value"
    ws["A2"], ws["B2"] = "alpha", 1
    ws["A3"], ws["B3"] = "beta", 2
    wb.save(HERE / "spreadsheet.xlsx")


def _make_pdf() -> None:
    """A minimal, hand-assembled single-page PDF (no reportlab/fpdf
    dependency) with two lines of real text content -- enough for
    docling_pdf-marked golden tests.
    """
    lines = ["Sample PDF Title", "This is a line of body text in the sample PDF."]
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        "/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    y = 750
    content_lines = ["BT", "/F1 18 Tf"]
    for line in lines:
        content_lines.append(f"1 0 0 1 72 {y} Tm ({line}) Tj")
        y -= 24
    content_lines.append("ET")
    content = "\n".join(content_lines)
    objects.append(f"<< /Length {len(content)} >>\nstream\n{content}\nendstream")

    pdf = "%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n{obj}\nendobj\n"
    xref_offset = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for off in offsets[1:]:
        pdf += f"{off:010} 00000 n \n"
    pdf += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    (HERE / "sample.pdf").write_bytes(pdf.encode("latin-1"))


if __name__ == "__main__":
    _make_docx()
    _make_pptx()
    _make_xlsx()
    _make_pdf()
    print("Fixtures written to", HERE)
